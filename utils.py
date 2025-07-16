import os
import csv
import json
import zipfile
import shutil
from pathlib import Path
from PyPDF2 import PdfReader
from docx import Document
from openpyxl import load_workbook
from PIL import Image
import datetime

SUPPORTED_EXTENSIONS = [
    ".pdf", ".docx", ".xlsx", ".json", ".csv", ".txt", ".jpg", ".jpeg", ".png", ".zip", ".log", ".xml"
]


def is_supported_file(file_path):
    return Path(file_path).suffix.lower() in SUPPORTED_EXTENSIONS


def is_mac_junk_file(file_path):
    name = Path(file_path).name
    return (
        name.startswith("__MACOSX") or 
        name.startswith("._") or 
        name.startswith(".DS_Store") or 
        name.startswith("Icon\r") or 
        name.startswith("_")
    )


def ensure_clean_dir(original_path):
    cleaned_folder = Path(original_path).parent / "cleaned"
    cleaned_folder.mkdir(exist_ok=True)
    cleaned_path = cleaned_folder / f"{Path(original_path).stem}_cleaned{Path(original_path).suffix}"
    return cleaned_path


def extract_metadata(file_path):
    ext = Path(file_path).suffix.lower()
    try:
        if ext == ".pdf":
            reader = PdfReader(file_path)
            meta = reader.metadata
            return {k[1:]: v for k, v in meta.items()} if meta else {}

        elif ext == ".docx":
            doc = Document(file_path)
            props = doc.core_properties
            return {
                "author": props.author,
                "comments": props.comments,
                "created": props.created,
                "last_modified_by": props.last_modified_by,
                "modified": props.modified,
                "revision": props.revision,
            }

        elif ext == ".xlsx":
            wb = load_workbook(file_path)
            props = wb.properties
            return {
                "author": props.creator,
                "title": props.title,
                "created": props.created,
                "last_modified_by": props.last_modified_by,
                "modified": props.modified,
            }

        elif ext == ".json":
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                data = json.load(f)
            return {"keys": list(data.keys()), "preview": str(data)[:300]}

        elif ext == ".csv":
            with open(file_path, newline='', encoding='utf-8', errors='ignore') as f:
                reader = csv.reader(f)
                headers = next(reader, [])
                return {"cabeceras": headers, "filas": sum(1 for _ in reader)}

        elif ext == ".txt" or ext == ".log":
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read(500)
            return {"preview": content}

        elif ext == ".xml":
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                lines = f.readlines()
            return {"lines": len(lines), "preview": "".join(lines[:10])}

        elif ext in [".jpg", ".jpeg", ".png"]:
            img = Image.open(file_path)
            return {"formato": img.format, "tamaño": img.size, "modo": img.mode}

        elif ext == ".zip":
            with zipfile.ZipFile(file_path, 'r') as z:
                return {"archivos_contenidos": z.namelist()}

    except Exception as e:
        return {"error": str(e)}

    return {}


def save_cleaned_file(original_path, cleaned_binary):
    out_path = ensure_clean_dir(original_path)
    with open(out_path, 'wb') as f:
        f.write(cleaned_binary)
    return out_path


def print_file_info(file_path):
    size_kb = os.path.getsize(file_path) / 1024
    print(f"📄 {Path(file_path).name} - {size_kb:.2f} KB - Tipo: {get_mime_type(file_path)}")


def get_mime_type(file_path):
    ext = Path(file_path).suffix.lower()
    mime_map = {
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".json": "application/json",
        ".csv": "text/csv",
        ".txt": "text/plain",
        ".log": "text/plain",
        ".xml": "application/xml",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".png": "image/png",
        ".zip": "application/zip",
    }
    return mime_map.get(ext, "unknown")


def get_supported_files(path):
    """Devuelve una lista de archivos válidos desde una ruta."""
    all_files = []
    if os.path.isdir(path):
        for root, _, files in os.walk(path):
            for f in files:
                full_path = os.path.join(root, f)
                if is_supported_file(full_path) and not is_mac_junk_file(full_path):
                    all_files.append(full_path)
    elif os.path.isfile(path) and is_supported_file(path) and not is_mac_junk_file(path):
        all_files.append(path)
    return all_files


def show_metadata_details(file_path, metadata):
    print(f"\n📂 Metadatos de: {file_path}")
    if not metadata:
        print("  ⚠️ No se encontraron metadatos o no se pudieron extraer.")
        return
    for k, v in metadata.items():
        print(f"  • {k}: {v}")